import os
import hashlib
import zipfile
import pypdf
import docx

def verify_submission():
    pkg_dir = "submission_package"
    zip_path = "JAI_Final_Submission_2026.zip"
    
    print("=========================================================================================================")
    print("FINAL SUBMISSION INTEGRITY AUDIT: JAI UMY SINTA 2 (SPRINT 3)")
    print("=========================================================================================================")
    
    files = sorted([f for f in os.listdir(pkg_dir) if not f.startswith('.') and not f.startswith('~$')])
    
    results = []
    all_clean = True
    
    for fname in files:
        fpath = os.path.join(pkg_dir, fname)
        size = os.path.getsize(fpath)
        
        # Calculate SHA256
        h = hashlib.sha256()
        with open(fpath, "rb") as f:
            while chunk := f.read(8192):
                h.update(chunk)
        sha256 = h.hexdigest()[:12] + "..."
        
        # Format check
        status = "PASS"
        detail = ""
        
        if size == 0:
            status = "FAIL (Zero-Byte)"
            all_clean = False
        elif fname.endswith(".docx"):
            try:
                d = docx.Document(fpath)
                detail = f"{len(d.paragraphs)} paras, {len(d.tables)} tables"
            except Exception as e:
                status = f"CORRUPT ({e})"
                all_clean = False
        elif fname.endswith(".pdf"):
            try:
                reader = pypdf.PdfReader(fpath)
                detail = f"{len(reader.pages)} pages"
            except Exception as e:
                status = f"CORRUPT ({e})"
                all_clean = False
        elif fname.endswith(".md"):
            detail = f"{len(open(fpath).readlines())} lines"
            
        results.append((fname, size, sha256, detail, status))
        
    print(f"\n{'File Name':<35} | {'Size':<12} | {'SHA-256':<16} | {'Content Summary':<25} | {'Status'}")
    print("-" * 105)
    for fname, size, sha256, detail, status in results:
        print(f"{fname:<35} | {size:>10,} B | {sha256:<16} | {detail:<25} | {status}")
        
    print("-" * 105)
    
    # Test zip archive
    print(f"\nAuditing Zip Archive: {zip_path} ({os.path.getsize(zip_path):,} bytes)...")
    with zipfile.ZipFile(zip_path, "r") as zf:
        bad_file = zf.testzip()
        if bad_file is None:
            print("  [PASS] CRC-32 Check: All files in zip archive passed without corruption.")
        else:
            print(f"  [FAIL] Corrupted file in zip: {bad_file}")
            all_clean = False
            
    print("\n=========================================================================================================")
    if all_clean:
        print("OVERALL AUDIT VERDICT: 100% PASS — READY FOR OJS SUBMISSION TO JAI UMY")
    else:
        print("OVERALL AUDIT VERDICT: FAIL — PLEASE REVIEW CORRUPT FILES")
    print("=========================================================================================================")

if __name__ == "__main__":
    verify_submission()
